#!/usr/bin/env python
# -*- coding: utf-8 -*-

"""Convertisseur d'un fichier PPTX et d'un fichier TXT en PowerPoint avec des notes.

On obtient en sortie le fichier avec les notes sous 3 formats différents : ODP, PPT et PPTX.

Usage : 
=======
    py PPTX_TXT_to_PPTX.py -i <input> -p <inputpptx> -o <output>

    <input> : chemin du fichier TXT avec les notes
    <inputpptx> : chemin du fichier PPTX à convertir
    <output> : nom du PowerPoint en sortie
"""

import getopt, sys
import collections
import collections.abc
from pptx import Presentation
import subprocess
import json
import base64
import os
import pipes
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
import config

def usage():
    """Explication de l'utilisation du programme.
    """
    
    print("Usage:")
    print("Minimal : ./PPTX_TXT_to_PPTX.py --input test_notes.txt --inputpptx test_pptx.pptx")
    print("Complete : ./PPTX_TXT_to_PPTX.py --input test_notes.txt --inputpptx test_pptx.pptx --output test_pptx_notes>")

def main():
    try:
        opts, args = getopt.getopt(sys.argv[1:], "hoip:v", ["help", "output=", "input=", "inputpptx="])
    except getopt.GetoptError as err:
        # print help information and exit:
        print(err)  # will print something like "option -a not recognized"
        usage()
        sys.exit(2)
    Translate = True
    WithNotes = False
    input=""
    output=""
    inputpptx=""
    lfrom="FR"
    lto="EN"
    C_ASPECTRATIO = "16/10"
    verbose = False
    for o, a in opts:
        if o == "-v":
            verbose = True
        elif o in ("-h", "--help"):
            usage()
            sys.exit()
        elif o in ("-i", "--input"):
            input = a
        elif o in ("-p", "--inputpptx"):
            inputpptx = a
        elif o in ("-o", "--output"):
            output = a
        else:
            assert False, "unhandled option"
    if Translate and input != "":
        if output == "":
            output="tmp/out.pptx"
        PPTX_TXTtoPPTX(input,inputpptx,output)
    else:
        usage()

def ExtractNotes(filename):
    """Extraction des notes pour chaque slide.
    
    Parameters
    ----------
    filename : str
        Fichier TXT dont on extrait les notes.

    Returns
    -------
    array
        Les notes et les slides associées.
    """

    if filename.split(".")[-1] == "pptx":
        prs = Presentation(filename)
        notes=[]
        num=0
        for slide in prs.slides:
            num+=1
            notesslide=[]
            for note in slide.notes_slide.notes_text_frame.paragraphs:
                notesslide.append(note.text)
            slidenotes=dict()
            slidenotes["Slide"]=num
            slidenotes["Notes"]=notesslide
            notes.append(slidenotes)
    else:
        content = open(filename, 'r', encoding="utf-8").read().split("\n")
        content.append('-----------')
        notes=[]
        num=0
        notesslide=[]
        for txt in content:
            if txt.find("----------")>=0 or txt.find("[forward]")>=0:
                if (num>0):
                    slidenotes=dict()
                    slidenotes["Slide"]=num
                    slidenotes["Notes"]=notesslide
                    notes.append(slidenotes)
                    notesslide=[]
                num+=1
            else:
                notesslide.append(txt)
    return notes

def PPTX_TXTtoPPTX(filename,filenamepptx,fileout):
    """Ajout des notes pour chaque slide.
    
    Parameters
    ----------
    filename : str
        Fichier TXT avec les notes.
    filenamepptx : str
        Fichier PPTX sans notes.
    fileout : str
        nom du PowerPoint en sortie
    """

    # Gestion des notes
    print("Extracting notes")
    Notes = ExtractNotes(filename)
    prs = Presentation(filenamepptx)
    nslide=0
    for slide in prs.slides:
        print("Adding notes to slide ",nslide+1)
        notes_slide = slide.notes_slide
        text_frame = notes_slide.notes_text_frame
        text_frame.clear()
        Notes.append({"Notes":[""]})
        text_frame.add_paragraph().text = "\n".join(Notes[nslide]["Notes"])
        nslide += 1
    prs.save(fileout+".pptx")
    outputdir=os.path.dirname(fileout)
    batcmd="HOME=/tmp "+config.SOFFICE+" --headless --convert-to odp "+fileout+".pptx"+" --outdir "+outputdir
    result = subprocess.check_output(batcmd, shell=True)
    batcmd="HOME=/tmp "+config.SOFFICE+" --headless --convert-to ppt "+fileout+".pptx"+" --outdir "+outputdir
    result = subprocess.check_output(batcmd, shell=True)


if __name__ == "__main__":
    main()

